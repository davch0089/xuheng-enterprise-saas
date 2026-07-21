"""ERP 基础资料 Excel 导入导出服务。"""

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel
from sqlalchemy import false, func, select

from apps.erp.master import models as master_models
from apps.erp.master import schemas as master_schemas
from apps.erp.master.crud import MasterDal
from core.exception import CustomException
from ..models import ErpDataExchangeTask
from ..schemas import ExchangeTaskOut


@dataclass(frozen=True)
class ExchangeField:
    """定义一个 Excel 列与模型字段的映射。"""

    key: str
    label: str
    required: bool = False
    value_type: str = "string"


@dataclass(frozen=True)
class ExchangeResource:
    """定义一个可导入导出的 ERP 资源。"""

    label: str
    model: type
    schema: type[BaseModel]
    fields: tuple[ExchangeField, ...]


BASE_FIELDS = (
    ExchangeField("code", "编码", True),
    ExchangeField("name", "名称", True),
    ExchangeField("is_active", "启用", value_type="boolean"),
    ExchangeField("remark", "备注"),
)
PARTNER_FIELDS = (
    ExchangeField("code", "编码", True), ExchangeField("name", "名称", True),
    ExchangeField("short_name", "简称"), ExchangeField("category", "分类"),
    ExchangeField("contact_name", "联系人"), ExchangeField("phone", "联系电话"),
    ExchangeField("email", "邮箱"), ExchangeField("tax_no", "税号"),
    ExchangeField("bank_name", "开户行"), ExchangeField("bank_account", "银行账号"),
    ExchangeField("address", "地址"), ExchangeField("payment_days", "账期天数", value_type="integer"),
    ExchangeField("is_active", "启用", value_type="boolean"), ExchangeField("remark", "备注"),
)


RESOURCES = {
    "product-categories": ExchangeResource(
        "商品分类", master_models.ErpProductCategory, master_schemas.ProductCategory,
        BASE_FIELDS[:2] + (ExchangeField("order", "排序", value_type="integer"),) + BASE_FIELDS[2:],
    ),
    "units": ExchangeResource(
        "计量单位", master_models.ErpUnit, master_schemas.Unit,
        BASE_FIELDS[:2] + (ExchangeField("symbol", "单位符号"), ExchangeField("decimal_places", "小数位", value_type="integer")) + BASE_FIELDS[2:],
    ),
    "warehouses": ExchangeResource(
        "仓库", master_models.ErpWarehouse, master_schemas.Warehouse,
        BASE_FIELDS[:2] + (
            ExchangeField("manager_name", "负责人"), ExchangeField("phone", "联系电话"),
            ExchangeField("address", "地址"), ExchangeField("allow_negative_stock", "允许负库存", value_type="boolean"),
        ) + BASE_FIELDS[2:],
    ),
    "customers": ExchangeResource(
        "客户", master_models.ErpCustomer, master_schemas.Customer,
        PARTNER_FIELDS[:-2] + (ExchangeField("credit_limit", "信用额度", value_type="decimal"),) + PARTNER_FIELDS[-2:],
    ),
    "suppliers": ExchangeResource("供应商", master_models.ErpSupplier, master_schemas.Supplier, PARTNER_FIELDS),
    "employees": ExchangeResource(
        "职员", master_models.ErpEmployee, master_schemas.Employee,
        BASE_FIELDS[:2] + (
            ExchangeField("department_id", "部门ID", value_type="integer"), ExchangeField("position", "职位"),
            ExchangeField("phone", "联系电话"), ExchangeField("email", "邮箱"),
        ) + BASE_FIELDS[2:],
    ),
    "settlement-methods": ExchangeResource(
        "结算方式", master_models.ErpSettlementMethod, master_schemas.SettlementMethod,
        BASE_FIELDS[:2] + (
            ExchangeField("method_type", "类型"), ExchangeField("payment_days", "默认账期", value_type="integer"),
        ) + BASE_FIELDS[2:],
    ),
}


class DataExchangeService:
    """生成 Excel 模板、导出基础资料并以编码幂等导入。"""

    def __init__(self, db, user=None):
        """绑定数据库会话和当前操作人。"""

        self.db = db
        self.user = user

    @staticmethod
    def resources() -> list[dict]:
        """返回前端可选择的数据资源。"""

        return [{"value": key, "label": value.label} for key, value in RESOURCES.items()]

    @staticmethod
    def config(resource: str) -> ExchangeResource:
        """获取资源配置，不支持时抛出业务异常。"""

        config = RESOURCES.get(resource)
        if not config:
            raise CustomException("不支持的数据导入导出类型")
        return config

    @staticmethod
    def workbook(config: ExchangeResource, rows: list[Any] | None = None) -> BytesIO:
        """生成带字段说明和示例的 Excel 文件流。"""

        book = Workbook()
        sheet = book.active
        sheet.title = config.label
        sheet.append([f"* {field.label}" if field.required else field.label for field in config.fields])
        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)
        sheet.freeze_panes = "A2"
        if rows is None:
            example = []
            for field in config.fields:
                if field.key == "code":
                    example.append("EXAMPLE001")
                elif field.key == "name":
                    example.append(f"示例{config.label}")
                elif field.value_type == "boolean":
                    example.append("是")
                elif field.value_type in {"integer", "decimal"}:
                    example.append(0)
                else:
                    example.append("")
            sheet.append(example)
        else:
            for row in rows:
                values = []
                for field in config.fields:
                    value = getattr(row, field.key, None)
                    if field.value_type == "boolean":
                        value = "是" if value else "否"
                    values.append(value)
                sheet.append(values)
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 4, 35)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        return stream

    async def template(self, resource: str) -> tuple[BytesIO, str]:
        """生成指定资源的空白导入模板。"""

        config = self.config(resource)
        return self.workbook(config), f"{config.label}导入模板.xlsx"

    async def export(self, resource: str) -> tuple[BytesIO, str]:
        """导出指定资源的全部未删除记录并登记任务。"""

        config = self.config(resource)
        rows = list(
            (
                await self.db.scalars(
                    select(config.model)
                    .where(config.model.is_delete == false())
                    .order_by(config.model.id)
                )
            ).all()
        )
        task = self._new_task("export", resource, f"{config.label}导出.xlsx")
        task.status = "success"
        task.total_count = len(rows)
        task.success_count = len(rows)
        self.db.add(task)
        await self.db.flush()
        return self.workbook(config, rows), f"{config.label}导出.xlsx"

    async def import_file(self, resource: str, file: UploadFile, mode: str = "upsert") -> dict:
        """读取 Excel，逐行校验并按编码新增或更新资料。"""

        if mode not in {"upsert", "create"}:
            raise CustomException("导入模式只能选择 upsert 或 create")
        config = self.config(resource)
        if not (file.filename or "").lower().endswith(".xlsx"):
            raise CustomException("只支持 .xlsx 文件")
        task = self._new_task("import", resource, file.filename)
        self.db.add(task)
        await self.db.flush()
        errors: list[dict] = []
        try:
            content = await file.read()
            if len(content) > 20 * 1024 * 1024:
                raise CustomException("导入文件不能超过 20MB")
            sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
            headers = [str(value or "").removeprefix("* ").strip() for value in next(sheet.iter_rows(values_only=True))]
            label_map = {field.label: field for field in config.fields}
            columns = [label_map.get(label) for label in headers]
            if not any(columns) or not all(field.label in headers for field in config.fields if field.required):
                raise CustomException("Excel 表头与模板不一致，请重新下载模板")
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            task.total_count = len([row for row in rows if any(value not in (None, "") for value in row)])
            for row_number, row in enumerate(rows, start=2):
                if not any(value not in (None, "") for value in row):
                    continue
                try:
                    values = {
                        field.key: self._convert(value, field.value_type)
                        for field, value in zip(columns, row)
                        if field and value not in (None, "")
                    }
                    validated = config.schema.model_validate(values).model_dump()
                    async with self.db.begin_nested():
                        existing = await self.db.scalar(
                            select(config.model).where(config.model.code == validated["code"])
                        )
                        if existing and mode == "create" and not existing.is_delete:
                            raise CustomException("编码已存在")
                        dal = MasterDal(self.db, config.model)
                        await dal.ensure_unique(validated, existing.id if existing else None)
                        if existing:
                            for key, value in validated.items():
                                setattr(existing, key, value)
                            existing.is_delete = False
                            existing.delete_datetime = None
                            self.db.add(existing)
                            await self.db.flush()
                        else:
                            await dal.create_data(validated)
                    task.success_count += 1
                except Exception as exc:
                    errors.append({"row": row_number, "message": self._error_message(exc)})
            task.failure_count = len(errors)
            task.error_details = errors[:200]
            task.status = "partial" if errors and task.success_count else "failed" if errors else "success"
        except Exception as exc:
            task.status = "failed"
            task.failure_count = task.total_count or 1
            task.error_details = [{"row": 0, "message": self._error_message(exc)}]
        self.db.add(task)
        await self.db.flush()
        await self.db.refresh(task)
        return ExchangeTaskOut.model_validate(task).model_dump()

    async def tasks(self, page: int, limit: int) -> tuple[list[dict], int]:
        """分页返回数据交换历史。"""

        statement = select(ErpDataExchangeTask).where(ErpDataExchangeTask.is_delete == false())
        count = await self.db.scalar(
            select(func.count(ErpDataExchangeTask.id)).where(
                ErpDataExchangeTask.is_delete == false()
            )
        )
        rows = (
            await self.db.scalars(
                statement.order_by(ErpDataExchangeTask.id.desc()).offset((page - 1) * limit).limit(limit)
            )
        ).all()
        return [ExchangeTaskOut.model_validate(row).model_dump() for row in rows], count or 0

    def _new_task(self, task_type: str, resource: str, file_name: str | None) -> ErpDataExchangeTask:
        """创建包含当前操作人的数据交换任务。"""

        return ErpDataExchangeTask(
            task_type=task_type,
            resource=resource,
            file_name=file_name,
            status="processing",
            operator_id=getattr(self.user, "id", None),
            operator_name=getattr(self.user, "name", None),
        )

    @staticmethod
    def _convert(value: Any, value_type: str) -> Any:
        """把 Excel 单元格转换为模型可验证的基础类型。"""

        if value_type == "boolean":
            if isinstance(value, bool):
                return value
            normalized = str(value).strip().lower()
            if normalized in {"是", "启用", "true", "1", "yes"}:
                return True
            if normalized in {"否", "停用", "false", "0", "no"}:
                return False
            raise ValueError(f"无法识别布尔值：{value}")
        if value_type == "integer":
            return int(value)
        if value_type == "decimal":
            return str(value)
        return str(value).strip() if value is not None else None

    @staticmethod
    def _error_message(exc: Exception) -> str:
        """提取业务异常文本，避免 CustomException 的默认字符串为空。"""

        return str(getattr(exc, "msg", None) or exc or "未知错误")[:500]
